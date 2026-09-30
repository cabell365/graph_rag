import hashlib
import json
import re
from pathlib import Path
import yaml

def digest(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as source:
        for block in iter(lambda: source.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()

def metadata(text):
    if text.startswith('---\n'):
        return yaml.safe_load(text.split('---', 2)[1]) or {}
    return {}

CHUNKER_VERSION = 'v2-2400-200'

def split_text(text, size=2400, overlap=200):
    # Conservative character bound keeps Nomic inputs comfortably under context.
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = text.rfind('\n', start + size // 2, end)
            if boundary > start:
                end = boundary
        yield text[start:end]
        if end == len(text):
            break
        start = max(start + 1, end - overlap)

def extract(path):
    if path.suffix == '.docx':
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        parts.extend(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        return '\n'.join(parts)
    if path.suffix == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for sheet in wb:
            parts.append(f'Sheet: {sheet.title}')
            parts.extend(' | '.join('' if v is None else str(v) for v in row)
                         for row in sheet.iter_rows(values_only=True))
        wb.close()
        return '\n'.join(parts)
    return path.read_text(encoding='utf-8-sig')

def documents(data):
    for path in sorted((data / 'corpus/documents').glob('*.md')):
        text = extract(path)
        meta = metadata(text)
        yield {'id': meta['document_id'], 'title': meta.get('title', path.stem),
               'path': str(path.relative_to(data)), 'text': text, 'sha256': digest(path),
               'origin': meta.get('fact_origin', 'unknown'),
               'effective_from': str(meta.get('effective_from', '')),
               'effective_to': str(meta.get('effective_to', '')),
               'status': meta.get('status', 'unknown'), 'entity_ids': meta.get('entity_ids', [])}
    for path in sorted(data.iterdir()):
        if path.suffix not in {'.md', '.docx', '.xlsx'}:
            continue
        text = extract(path)
        # Companion includes copies of the entire corpus. Keep its overview only;
        # all original business documents are already separately searchable.
        if path.name == 'Canyon_Office_GraphRAG_Documentation.md':
            text = text.split('\n## Extraction prompt')[0]
        yield {'id': 'SOURCE-' + hashlib.sha256(path.name.encode()).hexdigest()[:12],
               'title': path.name, 'path': path.name, 'text': text, 'sha256': digest(path),
               'origin': 'workbook_sample' if path.suffix == '.xlsx' else 'source_documentation',
               'effective_from': '', 'effective_to': '', 'status': 'reference', 'entity_ids': []}

def scan(source, data, report):
    inventory = []
    for path in sorted(source.rglob('*')):
        if path.is_file() and path.name != '.DS_Store':
            kind = 'reference_fixture' if '/graph/' in str(path) else 'document'
            if path.name == 'evaluation.jsonl':
                kind = 'held_out_evaluation'
            elif path.suffix == '.sql':
                kind = 'structured_finance'
            elif path.suffix in {'.yaml', '.yml'}:
                kind = 'configuration_not_retrieved'
            inventory.append({'path': str(path.relative_to(source)), 'bytes': path.stat().st_size,
                              'sha256': digest(path), 'role': kind})
    manifest = json.loads((data / 'corpus/manifest.json').read_text())
    for item in manifest:
        if digest(data / 'corpus' / item['path']) != item['sha256']:
            raise ValueError(f'Manifest hash mismatch: {item["path"]}')
    entities = json.loads((data / 'corpus/graph/entities.json').read_text())
    edges = json.loads((data / 'corpus/graph/relationships.json').read_text())
    entity_ids = {e['id'] for e in entities}
    if len(entity_ids) != len(entities):
        raise ValueError('Duplicate fixture entity IDs')
    for edge in edges:
        if edge['source_id'] not in entity_ids or edge['target_id'] not in entity_ids:
            raise ValueError('Fixture edge has an unknown endpoint')
    supplied_chunks = [json.loads(line) for line in (data / 'corpus/chunks.jsonl').read_text().splitlines()]
    for chunk in supplied_chunks:
        original = data / 'corpus/documents' / (chunk['document_id'] + '.md')
        if chunk['text'] != original.read_text():
            raise ValueError('Supplied chunk text differs from original document')
    import csv
    csv_counts = {}
    for name in ['entities','relationships']:
        with (data / f'corpus/graph/{name}.csv').open(newline='') as file:
            csv_counts[name] = sum(1 for _ in csv.DictReader(file))
    if csv_counts != {'entities':len(entities),'relationships':len(edges)}:
        raise ValueError('Fixture CSV and JSON row counts differ')
    docs = list(documents(data))
    result = {'files': inventory, 'manifest_verified': len(manifest),
              'fixture_entities':len(entities),'fixture_relationships':len(edges),
              'supplied_chunks_verified':len(supplied_chunks),'csv_rows':csv_counts,
              'documents': [{k: v for k, v in d.items() if k != 'text'} | {'characters': len(d['text'])}
                            for d in docs],
              'notes': ['SQL is parsed as data; never executed.',
                        'Graph JSON/CSV/Cypher are alternate representations, not separate evidence.',
                        'Evaluation answers are excluded from retrieval.',
                        'Original Compose password is excluded from the app and report.',
                        'Workbook is a separate sample, not a finance identity mapping.']}
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps(result, indent=2))
    return result
